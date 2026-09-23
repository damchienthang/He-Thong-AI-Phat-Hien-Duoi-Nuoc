"""Generate handoff DOCX and real browser screenshots with explicit test fixtures."""
import json
import zipfile
from pathlib import Path
from datetime import datetime
from playwright.sync_api import sync_playwright
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'ban_giao_3_2'
OUT.mkdir(exist_ok=True)
shots = []
with sync_playwright() as p:
    browser = p.chromium.launch(channel='msedge', headless=True)
    page = browser.new_page(viewport={'width': 1440, 'height': 1000})
    routes, messages, errors = [], [], []
    page.on('pageerror', lambda e: errors.append(str(e)))
    def route_ws(route):
        routes.append(route)
        route.on_message(lambda m: messages.append(json.loads(m)))
    page.route_web_socket('ws://localhost:8000/ws/monitor', route_ws)
    page.goto((ROOT / 'frontend/index.html').as_uri())
    page.evaluate("""() => {
      const b = document.createElement('div');
      b.textContent = 'KIỂM THỬ GIAO DIỆN • DỮ LIỆU GIẢ LẬP • KHÔNG PHẢI KẾT QUẢ AI';
      b.style.cssText = 'position:fixed;bottom:0;left:0;right:0;z-index:99999;background:#fff3cd;color:#493700;text-align:center;padding:7px;font:bold 14px Arial';
      document.body.appendChild(b);
    }""")
    page.locator('#file').set_input_files(str(ROOT / 'assets/demo/outdoor_pool.webm'))
    page.wait_for_function('video.readyState >= 2')
    page.evaluate('async () => {video.muted=true; await video.play();}')
    page.wait_for_function('video.currentTime > .5')
    page.evaluate('video.pause()')
    video_time = page.evaluate('video.currentTime')
    def capture(name, title, caption):
        page.evaluate("navigate('live')")
        page.wait_for_timeout(120)
        file = OUT / (name + '.png')
        page.screenshot(path=str(file), full_page=True)
        shots.append((file, title, caption))
    capture('01_cho_ket_noi', 'Chờ kết nối AI', 'Video phát độc lập; chưa có dự đoán nên không có box và số đếm hiển thị dấu —.')
    page.locator('[data-tab="settings"]').click()
    page.locator('#connect').click()
    page.wait_for_function("!document.getElementById('disconnect').disabled")
    frame = {'type':'frame','version':1,'stream_id':'kiem-thu-ui',
             'video_time':video_time,'model':'GIẢ LẬP KIỂM THỬ UI','is_demo':True,
             'persons':[{'id':1,'label':'swimming','confidence':.9,'duration':0,
                         'box':[.3,.35,.48,.75],'state':'NORMAL','event_id':'ui-01'}]}
    def send(state, duration=0, event='ui-01'):
        frame['persons'][0].update(state=state,duration=duration,event_id=event,
                                   label='swimming' if state=='NORMAL' else 'drowning')
        routes[0].send(json.dumps(frame))
        page.wait_for_function('(s) => latest && latest.persons[0].state === s', arg=state)
    send('NORMAL')
    capture('02_binh_thuong', 'Bình thường — khung xanh', 'Nhận NORMAL từ WebSocket giả lập: vẽ khung xanh và cập nhật một người đang theo dõi.')
    send('WARNING',10)
    capture('03_canh_bao', 'Cảnh báo — khung vàng', 'Nhận WARNING, duration = 10 giây từ dữ liệu kiểm thử. Backend chịu trách nhiệm xác định trạng thái.')
    send('CRITICAL',20)
    capture('04_nguy_hiem', 'Nguy hiểm — khung đỏ', 'Nhận CRITICAL, duration = 20 giây: khung đỏ và bảng cảnh báo khẩn. Đây không phải phát hiện đuối nước thật.')
    page.locator('#alerts .rescue').click()
    page.wait_for_timeout(100)
    assert messages[-1]['action'] == 'rescue_dispatched'
    routes[0].send(json.dumps({'type':'acknowledged','version':1,'stream_id':'kiem-thu-ui','event_id':'ui-01','person_id':1}))
    page.wait_for_function("document.querySelector('#alerts .rescue').textContent.includes('Backend')")
    capture('05_xac_nhan_cuu_ho', 'Xác nhận cứu hộ', 'Đã kiểm tra yêu cầu rescue_dispatched và phản hồi ACK giả lập. Xác nhận triển khai cứu hộ không tự đổi nguy hiểm thành an toàn.')
    send('CRITICAL',20,'ui-02')
    page.wait_for_function("!document.querySelector('#alerts .false-alarm').disabled")
    page.locator('#alerts .false-alarm').click()
    page.wait_for_timeout(100)
    assert messages[-1]['action'] == 'false_alarm'
    routes[0].send(json.dumps({'type':'acknowledged','version':1,'stream_id':'kiem-thu-ui','event_id':'ui-02','person_id':1,'action':'false_alarm'}))
    page.wait_for_function("document.querySelector('#alerts .rescue').textContent.includes('báo giả')")
    capture('06_bao_dong_gia', 'Xác nhận báo động giả', 'Đã kiểm tra yêu cầu false_alarm và ACK đúng sự cố. Việc lưu cơ sở dữ liệu cần backend thật.')
    assert not errors, errors
    browser.close()

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)
sec.top_margin = sec.bottom_margin = Inches(.55)
sec.left_margin = sec.right_margin = Inches(.65)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Arial', Pt(10)
normal.paragraph_format.space_after = Pt(5)
for name in ['Title','Heading 1','Heading 2']:
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = RGBColor.from_string('17365D')
doc.add_heading('BÁO CÁO BÀN GIAO', 0)
doc.add_paragraph('Nhiệm vụ 3.2 • Web Dashboard giám sát cứu hộ', style='Subtitle')
doc.add_paragraph('Ngày lập: ' + datetime.now().strftime('%d/%m/%Y'))
doc.add_heading('1. Kết quả thực hiện', 1)
doc.add_paragraph('Đã xây dựng dashboard theo mẫu giao diện: hiển thị video, khung trạng thái xanh/vàng/đỏ, bảng cảnh báo khẩn, nhật ký sự kiện và các nút xác nhận cứu hộ, báo động giả. Có WebSocket client để nhận dữ liệu và gửi thao tác xử lý. Mã giao diện ở frontend/index.html; app.py dùng Streamlit để mở giao diện.')
doc.add_heading('2. Kiểm thử và minh chứng', 1)
doc.add_paragraph('Đã chạy kiểm thử trình duyệt Edge với WebSocket giả lập: nhận NORMAL/WARNING/CRITICAL, vẽ box, cập nhật số đếm, gửi hai thao tác và nhận ACK. Sáu ảnh ở các trang sau chụp trực tiếp giao diện. Box được đặt bằng dữ liệu kiểm thử, không phải kết quả mô hình trên video.')
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Shading Accent 1'
table.rows[0].cells[0].text = 'Đầu ra theo phân công'
table.rows[0].cells[1].text = 'Tình trạng bàn giao'
for a,b in [('Mã nguồn Web Dashboard','Đã có giao diện và WebSocket client.'),('Ảnh trạng thái hoạt động','6 ảnh kiểm thử kèm trong tài liệu.'),('Khung màu, cảnh báo và thao tác cứu hộ','Đã kiểm thử bằng dữ liệu giả lập.'),('Tích hợp AI / FastAPI / cơ sở dữ liệu','Chờ backend và mô hình thật; chưa nghiệm thu toàn hệ thống.')]:
    cells=table.add_row().cells
    cells[0].text,cells[1].text=a,b
doc.add_heading('3. Nội dung cần phối hợp tiếp', 1)
doc.add_paragraph('Thống nhất URL và schema WebSocket với bên backend; đồng bộ nguồn video, stream_id và video_time; nhận box, trạng thái và ACK từ dịch vụ thật. Backend phụ trách AI, ngưỡng cảnh báo 10/20 giây và lưu sự cố. Sau tích hợp cần kiểm thử video thực tế, mất kết nối, độ trễ và âm thanh trên máy sử dụng.')
doc.add_heading('4. Cách chạy và tài liệu', 1)
doc.add_paragraph('Tại thư mục cam_ai, chạy trong PowerShell:\n.\\.venv\\Scripts\\python.exe -m streamlit run app.py\nHướng dẫn cài đặt: README.md. Hợp đồng tích hợp: docs/HANDOFF.md. Quy trình kiểm thử: docs/TESTING.md.')
doc.add_paragraph('Lưu ý: mở video local chỉ phát video trên trình duyệt, chưa gửi video lên AI. Báo cáo này chỉ bàn giao giao diện thuộc nhiệm vụ 3.2; chưa xác nhận độ chính xác mô hình hoặc hiệu năng hệ thống tích hợp.')
for i,(file,title,caption) in enumerate(shots):
    if i % 2 == 0:
        doc.add_page_break()
        doc.add_heading('Ảnh minh chứng kiểm thử giao diện', 1)
    doc.add_paragraph(f'Hình {i+1}. {title}', style='Heading 2')
    doc.add_picture(str(file), width=Inches(5.75))
    para=doc.add_paragraph(caption)
    para.paragraph_format.space_after=Pt(8)
    for run in para.runs: run.font.size=Pt(9)
footer=sec.footer.paragraphs[0]
footer.text='Nhiệm vụ 3.2 • Ảnh dùng dữ liệu giả lập, chưa tích hợp AI thật'
footer.style=doc.styles['Caption']
doc.core_properties.title='Bàn giao Web Dashboard giám sát cứu hộ — Nhiệm vụ 3.2'
doc.core_properties.author=''
doc.save(OUT/'Bao_cao_ban_giao_Dashboard_3_2.docx')
(OUT/'NGUON_ANH.txt').write_text('Ảnh chụp trực tiếp frontend; toàn bộ box/cảnh báo là dữ liệu WebSocket giả lập.\nVideo nền: Outdoor pool.webm — Almanta, CC BY-SA 4.0.\nhttps://commons.wikimedia.org/wiki/File:Outdoor_pool.webm\nhttps://creativecommons.org/licenses/by-sa/4.0/\nCác ảnh dẫn xuất từ video được chia sẻ theo CC BY-SA 4.0. Không ngụ ý tác giả xác nhận ứng dụng.\n',encoding='utf-8')
# Include attribution inside the Word report as well.
doc.add_paragraph('Nguồn video nền: Outdoor pool.webm — Almanta (Wikimedia Commons), CC BY-SA 4.0. Các ảnh dẫn xuất theo cùng giấy phép.\nhttps://commons.wikimedia.org/wiki/File:Outdoor_pool.webm\nhttps://creativecommons.org/licenses/by-sa/4.0/', style='Caption')
doc.save(OUT/'Bao_cao_ban_giao_Dashboard_3_2.docx')
with zipfile.ZipFile(OUT/'Anh_trang_thai_3_2.zip','w',zipfile.ZIP_DEFLATED) as z:
    for file,_,_ in shots: z.write(file,file.name)
    z.write(OUT/'NGUON_ANH.txt','NGUON_ANH.txt')
print('PASS: 6 screenshots, green/yellow/red states, rescue and false-alarm ACK; no browser errors.')
print(OUT/'Bao_cao_ban_giao_Dashboard_3_2.docx')
